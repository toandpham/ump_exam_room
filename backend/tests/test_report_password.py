"""Báo cáo có mật khẩu: mở được bằng Excel, KHÔNG cần giải nén (AD-123).

Bản cũ bọc file trong ZIP mã hoá AES-256 của WinZip. Trường báo "giải nén không
được" — đúng: **Windows Explorer không mở được kiểu mã hoá đó** (nó chỉ hiểu
ZipCrypto cũ), phải cài 7-Zip mới mở nổi. Nay đặt mật khẩu NGAY TRONG file Excel
theo chuẩn Office: double-click, Excel hỏi mật khẩu, xong.

Dùng dữ liệu báo cáo THẬT (qua build_report) chứ không tự chế fixture — fixture tự
chế sẽ lệch khỏi cấu trúc thật rồi test mất giá trị.
"""

import io

import msoffcrypto
import openpyxl
import pytest

from app.services import report_service
from tests.conftest import auth, fast_forward_start

pytestmark = pytest.mark.asyncio

PW = "matkhau123"


async def _real_report(client, factory, db):
    admin, ptok = await factory.admin()
    exam, sitting, payload = await factory.active_exam(
        [{"text": "Q1", "correct": "A"}], owner_id=admin.id)
    cand = await factory.candidate(exam.id)
    tok = (await client.post("/api/exam/auth/login",
                             json={"cccd": cand.cccd})).json()["token"]
    ch = auth(tok)
    await client.post("/api/exam/auth/confirm", headers=ch)
    await client.post(f"/api/admin/sittings/{sitting.id}/start", headers=auth(ptok))
    await fast_forward_start(sitting.id)
    await client.post("/api/exam/answer",
                      json={"question_id": payload["questions"][0]["id"],
                            "selected_option": "A"}, headers=ch)
    await client.post("/api/exam/submit", headers=ch)
    report = await report_service.build_report(db, sitting, exam.name)
    return report, cand


def _decrypt(data: bytes, password: str) -> io.BytesIO:
    out = io.BytesIO()
    f = msoffcrypto.OfficeFile(io.BytesIO(data))
    f.load_key(password=password)
    f.decrypt(out)
    out.seek(0)
    return out


async def test_encrypted_xlsx_opens_with_the_password(client, factory, db):
    report, cand = await _real_report(client, factory, db)
    data = report_service.export_excel_encrypted(report, PW)

    # Vỏ ngoài phải là OLE2 — đó là thứ Excel nhận ra để hỏi mật khẩu. Nếu ra "PK"
    # thì lại là zip, tức quay về đúng lỗi cũ.
    assert data[:4] == b"\xd0\xcf\x11\xe0", data[:8].hex()

    wb = openpyxl.load_workbook(_decrypt(data, PW))
    assert wb.sheetnames == ["Kết quả", "Đáp án"], wb.sheetnames
    cccds = [c.value for c in wb["Kết quả"]["B"]]
    assert cand.cccd in cccds, cccds


async def test_wrong_password_is_rejected(client, factory, db):
    report, _ = await _real_report(client, factory, db)
    data = report_service.export_excel_encrypted(report, PW)
    with pytest.raises(Exception):
        _decrypt(data, "sai-mat-khau")


async def test_content_not_readable_without_password(client, factory, db):
    """Không có mật khẩu thì không đọc trộm được nội dung — nếu chỉ khoá sửa mà vẫn
    đọc được tên + điểm thì việc đặt mật khẩu chỉ là hình thức."""
    report, cand = await _real_report(client, factory, db)
    data = report_service.export_excel_encrypted(report, PW)
    with pytest.raises(Exception):
        openpyxl.load_workbook(io.BytesIO(data))
    assert cand.cccd.encode() not in data     # CCCD không nằm dạng rõ trong file
